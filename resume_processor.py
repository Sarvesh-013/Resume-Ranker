from kafka import KafkaConsumer
import PyPDF2
import docx
import os
import json
import re
import mysql.connector
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
import logging
from tqdm import tqdm
import sys
import socketio
import requests
import time
import ollama

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    filename='resume_processor.log',
    format='%(asctime)s %(levelname)s:%(message)s'
)

# Kafka Configuration
KAFKA_TOPIC = "resume-topic"
KAFKA_SERVER = "localhost:9092"

# SocketIO client
sio = socketio.Client()

try:
    sio.connect('http://localhost:5000')
    logging.info("Connected to SocketIO server")
except Exception as e:
    logging.error(f"Failed to connect to SocketIO server: {e}")

try:
    consumer = KafkaConsumer(
        KAFKA_TOPIC,
        bootstrap_servers=KAFKA_SERVER,
        auto_offset_reset="earliest",
        group_id="resume-consumer-group",
        value_deserializer=lambda x: json.loads(x.decode("utf-8")),
        enable_auto_commit=False
    )
    logging.info(f"Connected to Kafka broker, subscribed to topic: {KAFKA_TOPIC}")
except Exception as e:
    logging.error(f"Failed to initialize Kafka consumer: {e}")
    raise

def extract_text(file_path):
    """Extract text from a resume (PDF/DOCX/TXT)."""
    try:
        if file_path.endswith(".pdf"):
            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
                return text.strip()

        elif file_path.endswith(".docx"):
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            return text.strip()

        elif file_path.endswith(".txt"):
            with open(file_path, "r", encoding="utf-8") as file:
                return file.read().strip()

        else:
            logging.error(f"Unsupported file format: {file_path}")
            return None

    except Exception as e:
        logging.error(f"Error extracting text from {file_path}: {e}")
        return None

def extract_json(response_text):
    """Extract and validate JSON from response text."""
    # Log raw response for debugging
    logging.info(f"Raw API response: {response_text[:1000]}")  # Limit to avoid huge logs
    brace_count = 0
    start = -1
    for i, char in enumerate(response_text):
        if char == '{':
            if start == -1:
                start = i
            brace_count += 1
        elif char == '}':
            brace_count -= 1
            if brace_count == 0 and start != -1:
                json_text = response_text[start:i + 1]
                try:
                    parsed = json.loads(json_text)
                    # Validate required fields
                    required_fields = ["name", "phone", "email", "skills", "summary", "job_score", "positives"]
                    if all(field in parsed for field in required_fields):
                        return parsed
                    else:
                        return {
                            "error": f"Missing required fields: {', '.join(f for f in required_fields if f not in parsed)}",
                            "raw_output": response_text,
                            "extracted_text": json_text
                        }
                except json.JSONDecodeError as e:
                    return {
                        "error": f"Failed to parse extracted JSON: {e}",
                        "raw_output": response_text,
                        "extracted_text": json_text
                    }
    if start != -1 and brace_count > 0:
        json_text = response_text[start:] + '}'
        logging.info(f"Attempted fix by adding closing brace: {json_text[:1000]}")
        try:
            parsed = json.loads(json_text)
            required_fields = ["name", "phone", "email", "skills", "summary", "job_score", "positives"]
            if all(field in parsed for field in required_fields):
                return parsed
            else:
                return {
                    "error": f"Missing required fields: {', '.join(f for f in required_fields if f not in parsed)}",
                    "raw_output": response_text,
                    "attempted_text": json_text
                }
        except json.JSONDecodeError as e:
            return {
                "error": f"Failed to parse after fix attempt: {e}",
                "raw_output": response_text,
                "attempted_text": json_text
            }
    return {
        "error": "No complete JSON object found",
        "raw_output": response_text
    }

def process_resume_with_llama(resume_text, job_description, required_skills):
    """Extracts structured data and assigns a relevance score using remote Llama 3 API."""
    prompt = f"""
Extract the relevant details from the given resume and return a valid JSON object.

### *STRICT RULES:*
- *Output ONLY JSON.* No explanations, reasoning, or formatting notes.
- *Ensure valid JSON.* No missing brackets, extra characters, or incorrect structures.
- *Follow the exact format* without placeholders.
- If any data is missing, explicitly mention "None" in the respective field.
- *Phone, Email, Skills* must be strings, not lists.

### *Data Extraction Requirements:*
- *name:* Extract the full name accurately.
- *phone:* Extract a valid phone number as a string (or "None" if missing).
- *email:* Extract a valid email address as a string (or "None" if missing).
- *Skills:* Extract all skills which are only listed in the resume, not the ones present in the input and infer additional skills from project descriptions.
- *summary:* Provide a brief professional summary if present, or "None."
- *job_score:* Score the applicant based on job relevance (0-100) considering his skills and projects related to the mentioned job category.
- *positives:* Highlight key strengths as a string.

---

### *Scoring Guidelines (0-100 scale):*
- *Job Relevance (30%)* → Does the role & experience match the job description?
- *Skills (50%)* → Are the required skills explicitly mentioned or inferred?
- *Projects & Experience (20%)* → Are past projects/work experience aligned with the job?

*Scoring Deductions:*
- Missing must-have skills → Major deduction.
- Unrelated experience → Moderate deduction.
- Lack of depth in resume content → Minor deduction.

*Bonus Points:*
- Strong project work in the relevant field.
- Additional valuable skills beyond the required ones.
- Industry certifications, publications, or notable achievements.

---

### *Required Skills:* {required_skills}
### *Job Description:* {job_description}

Resume Content:
{resume_text}

### *Expected JSON Output:*
{{
  "name": "<Full Name or 'None'>",
  "phone": "<Phone Number or 'None'>",
  "email": "<Email or 'None'>",
  "skills": "<Comma-separated skills or 'None'>",
  "summary": "<Brief summary or 'None'>",
  "job_score": <Score between 0-100>,
  "positives": "<Briefly highlight strengths>"
}}
    """

    url = "https://tasty-pandas-cough.loca.lt/api/generate"
    payload = {
        "model": "llama3.1:8b",
        "prompt": prompt
    }
    max_retries = 3
    retry_delay = 5

    for attempt in range(max_retries):
        try:
            logging.info(f"Sending request to API (attempt {attempt + 1}/{max_retries})")
            response = requests.post(
                url,
                json=payload,
                headers={
                    "Content-Type": "application/json",
                    "bypass-tunnel-reminder": "true"
                },
                timeout=60,
                stream=True
            )
            response.raise_for_status()

            # Handle NDJSON response
            raw_output = ""
            for line in response.iter_lines(decode_unicode=True):
                if line:  # Skip empty lines
                    try:
                        json_line = json.loads(line)
                        if "response" in json_line:
                            raw_output += json_line["response"]
                        if json_line.get("done", False):
                            break  # Stop when done is true
                    except json.JSONDecodeError as e:
                        logging.error(f"Failed to parse NDJSON line: {line}, Error: {e}")
                        continue

            if not raw_output:
                logging.error("Empty response from API")
                return {"error": "Empty response from API", "raw_output": ""}

            logging.info("Received response from API")
            extracted_data = extract_json(raw_output)
            return extracted_data

        except requests.exceptions.RequestException as e:
            logging.error(f"API error (attempt {attempt + 1}): {str(e)}, Response: {e.response.text if e.response else 'No response'}")
            if attempt < max_retries - 1:
                logging.info(f"Retrying after {retry_delay} seconds...")
                time.sleep(retry_delay)
            else:
                logging.error("Max retries reached. Falling back to local Ollama.")
                try:
                    logging.info("Sending request to local Ollama server")
                    ollama_response = ollama.chat(
                        model="llama3.2",
                        messages=[{"role": "user", "content": prompt}]
                    )
                    raw_output = ollama_response["message"]["content"]
                    logging.info("Received response from Ollama")
                    return extract_json(raw_output)
                except Exception as e2:
                    logging.error(f"Local Ollama error: {str(e2)}")
                    return {"error": f"API error: {str(e)}, Local error: {str(e2)}", "raw_output": ""}
        except ValueError as e:
            logging.error(f"Response parsing error: {str(e)}")
            return {"error": f"Invalid response: {str(e)}", "raw_output": ""}

def save_to_database(resume_data, file_path):
    try:
        conn = mysql.connector.connect(
            host="localhost",
            user="root",
            password="sarvesh#13",
            database="resume_db"
        )
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM resumes WHERE file_path = %s", (file_path,))
        if cursor.fetchone():
            logging.info(f"Skipped duplicate resume: {file_path}")
            sio.emit('resume_processed', {
                'file_path': file_path,
                'status': 'failed',
                'error': 'Duplicate resume'
            })
            return

        query = """
        INSERT INTO resumes (name, phone, email, skills, summary, job_score, positives, uploaded_at, file_path)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """
        values = (
            resume_data.get("name", "None"),
            resume_data.get("phone", "None"),
            resume_data.get("email", "None"),
            resume_data.get("skills", "None"),
            resume_data.get("summary", "None"),
            resume_data.get("job_score", 0),
            resume_data.get("positives", "None"),
            datetime.now(),
            file_path
        )

        cursor.execute(query, values)
        conn.commit()
        logging.info(f"Successfully saved resume: {file_path}")
        cursor.close()
        conn.close()

        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'completed',
            'data': resume_data
        })

    except mysql.connector.Error as e:
        logging.error(f"Database error while saving resume: {file_path}, Error: {e}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': str(e)
        })

def process_message(message, consumer):
    logging.info(f"Received Kafka message: {message.value}")
    message_data = message.value
    file_path = message_data.get("file_path", "")
    job_description = message_data.get("job_description", "default job description")
    required_skills = message_data.get("skills", [])

    if not os.path.exists(file_path):
        logging.error(f"File not found: {file_path}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': 'File not found'
        })
        try:
            consumer.commit()
            logging.info(f"Committed Kafka offset for message: {message.offset}")
        except Exception as e:
            logging.error(f"Failed to commit Kafka offset: {e}")
        return

    logging.info(f"Started processing: {file_path}")
    sio.emit('resume_processing', {
        'file_path': file_path,
        'status': 'started'
    })

    resume_text = extract_text(file_path)

    if resume_text:
        extracted_details = process_resume_with_llama(resume_text, job_description, required_skills)
        logging.info(f"Extracted details for: {file_path}")
        save_to_database(extracted_details, file_path)
    else:
        logging.error(f"Failed to extract text from: {file_path}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': 'Failed to extract text'
        })

    try:
        consumer.commit()
        logging.info(f"Committed Kafka offset for message: {message.offset}")
    except Exception as e:
        logging.error(f"Failed to commit Kafka offset: {e}")

logging.info("Starting Kafka consumer loop")
with ThreadPoolExecutor(max_workers=3) as executor:
    progress_bar = tqdm(desc="Processed Resumes", unit="resume", file=sys.stdout)
    futures = []
    batch_size = 3  # Match max_workers
    for message in consumer:
        logging.info("Fetched a message from Kafka")
        futures.append(executor.submit(process_message, message, consumer))
        progress_bar.update(1)
        if len(futures) >= batch_size:
            for future in futures:
                future.result()  # Wait for threads to complete
            futures = []  # Clear completed futures
    # Process any remaining futures
    for future in futures:
        future.result()
logging.info("Consumer loop ended")